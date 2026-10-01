"""Dashboard TAT uses the same milestones as detail views and respects scope."""

from datetime import UTC, datetime, timedelta
from io import BytesIO

import pymupdf
import pytest
from openpyxl import load_workbook

from common.enums import ApplicationStatus as AS
from common.enums import UserRole
from dashboard import exports, pdf, services
from dashboard.services import DetailQuery, Scope
from pipeline.models import Application, Onboarding
from pipeline.services.tat import application_tat
from tests.test_jobs_support import make_application, make_jd
from tests.test_recruitment_tat import event

pytestmark = pytest.mark.django_db
START = datetime(2026, 9, 1, tzinfo=UTC)


def day(n):
    return START + timedelta(days=n)


@pytest.fixture
def owner(user_factory):
    return user_factory(role=UserRole.HR_ADMIN, timezone="UTC")


@pytest.fixture
def job(owner):
    return make_jd(owner, published_at=START)


def hire(job, *, added=2, completed=12, history=True):
    app = make_application(job, AS.ONBOARDED, stage_entered_at=day(completed))
    Application.objects.filter(pk=app.pk).update(created_at=day(added))
    app.refresh_from_db()
    if history:
        event(app, completed, AS.NEW, AS.ONBOARDED)
    return app


def scope(owner, **kwargs):
    return Scope(viewer=owner, start=day(10).date(), end=day(20).date(), now=day(30), **kwargs)


def test_averages_use_completion_window_and_previous_window(owner, job):
    hire(job, added=1, completed=8)
    first = hire(job, completed=12)
    hire(job, completed=18)
    hire(job, completed=31)  # Future completion is excluded.
    make_application(job, AS.ONBOARDING, stage_entered_at=day(16))
    result = services.summary(scope(owner))
    assert result["recruitment_tat"]["value"] == 15
    assert result["recruitment_tat"]["delta"] == 7
    assert result["time_to_hire"]["value"] == 13
    assert result["time_to_hire"]["delta"] == 6
    assert "2 hires" in result["recruitment_tat"]["detail"]
    details = services.details(scope(owner), DetailQuery(metric="recruitment_tat"))
    assert details["count"] == 2
    first_row = next(item for item in details["items"] if item["id"] == str(first.pk))
    expected = application_tat(first, now=day(30))["job_elapsed_seconds"] / 86400
    assert first_row["value"].startswith(f"{expected:.1f} days")
    assert first_row["href"].endswith("&tab=timeline")


def test_onboarding_completion_wins_over_stage_timestamp(owner, job):
    app = hire(job, completed=25)
    Onboarding.objects.create(application=app, start_date=day(10).date(), completed_at=day(14))
    result = services.summary(scope(owner))
    assert result["recruitment_tat"]["value"] == 14
    assert result["time_to_hire"]["value"] == 12


def test_missing_or_invalid_publication_excluded_from_job_tat_only(owner):
    for published in [None, day(20)]:
        hire(make_jd(owner, published_at=published), completed=12)
    result = services.summary(scope(owner))
    assert result["recruitment_tat"]["value"] is None
    assert result["time_to_hire"]["value"] == 10
    assert services.details(scope(owner), DetailQuery(metric="recruitment_tat"))["count"] == 0
    assert services.details(scope(owner), DetailQuery(metric="time_to_hire"))["count"] == 2


def test_people_filter_applies_to_totals_stages_and_details(owner, job, user_factory):
    other = user_factory(role=UserRole.HR)
    hire(job, completed=12)
    hire(make_jd(other, published_at=day(5)), completed=18)
    selected = scope(owner, user_ids=(str(other.pk),))
    assert services.summary(selected)["recruitment_tat"]["value"] == 13
    assert services.recruitment_stage_tat(selected)["hires"] == 1
    assert services.details(selected, DetailQuery(metric="recruitment_tat"))["count"] == 1


def test_stage_averages_sum_repeat_visits_and_exclude_incomplete_history(owner, job):
    app = hire(job, completed=12, history=False)
    event(app, 3, AS.NEW, AS.HR_REVIEW)
    event(app, 5, AS.HR_REVIEW, AS.ON_HOLD, bulk=True)
    event(app, 7, AS.ON_HOLD, AS.HR_REVIEW)
    event(app, 12, AS.HR_REVIEW, AS.ONBOARDED)
    second = hire(job, completed=14, history=False)
    event(second, 4, AS.NEW, AS.HR_REVIEW)
    event(second, 14, AS.HR_REVIEW, AS.ONBOARDED)
    hire(job, completed=16, history=False)
    result = services.recruitment_stage_tat(scope(owner))
    assert result["hires"] == 3
    assert result["incomplete_histories"] == 1
    stages = {row["key"]: row for row in result["stages"]}
    assert stages[AS.HR_REVIEW]["avg_days"] == 8.5  # (2 + 5 + 10) / 2 hires
    assert stages[AS.HR_REVIEW]["hires"] == 2
    assert stages[AS.ON_HOLD]["avg_days"] == 2
    assert AS.ONBOARDED not in stages


def test_history_query_count_does_not_grow_per_hire(owner, job, django_assert_num_queries):
    for completed in range(11, 20):
        hire(job, completed=completed)
    with django_assert_num_queries(2):
        assert services.recruitment_stage_tat(scope(owner))["hires"] == 9


def test_empty_period_has_unknown_totals_and_no_stage_bars(owner):
    result = services.summary(scope(owner))
    assert result["recruitment_tat"]["value"] is None
    assert result["recruitment_tat"]["delta"] is None
    assert result["time_to_hire"]["value"] is None
    assert services.recruitment_stage_tat(scope(owner)) == {
        "hires": 0,
        "incomplete_histories": 0,
        "stages": [],
    }


def test_date_boundaries_use_viewer_timezone(owner, job):
    owner.timezone = "Asia/Kolkata"
    hire(job, completed=9 + 23 / 24)
    selected = Scope(viewer=owner, start=day(10).date(), end=day(10).date(), now=day(30))
    assert services.recruitment_stage_tat(selected)["hires"] == 1
    hire(job, completed=10 + 20 / 24)  # Following local day.
    assert services.recruitment_stage_tat(selected)["hires"] == 1


def test_exports_include_recruitment_candidate_and_stage_tat(owner, job):
    hire(job)
    snap = exports.snapshot(scope(owner))
    csv = exports.to_csv(snap).decode("utf-8-sig")
    assert "Recruitment TAT" in csv and "Candidate TAT" in csv and "Stage TAT" in csv
    workbook = load_workbook(BytesIO(exports.to_xlsx(snap)))
    assert "Stage TAT" in workbook.sheetnames
    with pymupdf.open(stream=pdf.render(snap), filetype="pdf") as document:
        text = "\n".join(page.get_text() for page in document)
    assert "RECRUITMENT TAT" in text
    assert "Candidate TAT" in text and "Hires visiting stage" in text


def test_api_shapes_and_hr_only_access(owner, job, api_client, user_factory):
    hire(job)
    api_client.force_authenticate(owner)
    params = {"start": day(10).date(), "end": day(20).date()}
    assert "recruitment_tat" in api_client.get("/api/v1/dashboard/summary/", params).data
    assert "tat" in api_client.get("/api/v1/dashboard/insights/", params).data
    details = api_client.get("/api/v1/dashboard/details/", {**params, "metric": "recruitment_tat"})
    assert details.status_code == 200
    assert details.data["count"] == 1
    api_client.force_authenticate(user_factory(role=UserRole.INTERVIEWER))
    for endpoint in ["summary", "insights", "details"]:
        response = api_client.get(
            f"/api/v1/dashboard/{endpoint}/", {**params, "metric": "recruitment_tat"}
        )
        assert response.status_code == 403
