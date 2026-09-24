"""The dashboard as a file: every figure and table the page shows, for the same
window and people, as one CSV, one Excel workbook or one PDF."""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from html import escape
from typing import Any

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.models import User
from common import enums
from dashboard import services
from dashboard.services import Scope

Cell = str | int | float | date | datetime | None


@dataclass(frozen=True)
class Table:
    title: str
    columns: list[str]
    rows: list[list[Cell]]
    # One line under the title: what "now" means, the role the funnel is narrowed to…
    note: str = ""


@dataclass(frozen=True)
class Report:
    title: str
    # The window, the people, and who exported it when.
    lines: list[str]
    tables: list[Table]


ATTENTION_ITEMS = (
    ("overdue_follow_ups", "Overdue follow-ups", "Contacted candidates, next action past due"),
    ("feedback_pending", "Interview feedback owed", "Held, but no feedback submitted yet"),
    ("offers_expiring", "Offers expiring", "Past expiry or due within 3 days"),
    ("stale_candidates", "Stuck for a week", "No stage change in 7 days"),
    ("quiet_roles", "Quiet roles", "Nothing logged in 7 days"),
)
HEADLINE_FIGURES = (
    ("open_roles", "Open roles", "now"),
    ("in_pipeline", "In pipeline", "now"),
    ("new_candidates", "New candidates", "window"),
    ("interviews", "Interviews", "window"),
    ("offers_pending", "Offers pending", "now"),
    ("hires", "Hires", "window"),
    ("offer_acceptance", "Offer acceptance", "window"),
    ("time_to_hire", "Time to hire", "window"),
)
TREND_SERIES = (
    ("candidates", "Candidates found"),
    ("shortlisted", "Shortlisted"),
    ("interviews", "Interviews"),
    ("offers", "Offers sent"),
    ("hires", "Hires"),
)
ROLE_STAGES = ("shortlisted", "contacted", "interviewed", "selected", "onboarded")
ACTIVE_ROLE_STATUSES = (enums.JDStatus.OPEN, enums.JDStatus.ON_HOLD)
WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
UNITS = {"count": "", "percent": "%", "days": " days"}
DELTA_UNITS = {"count": "", "percent": " pp", "days": " days"}


def window_label(scope: Scope) -> str:
    """Last 30 days (26 Aug to 24 Sep 2026); a custom window is just its dates."""
    dates = scope.dates
    span = f"{dates[0]:%d %b %Y} to {dates[-1]:%d %b %Y}"
    return span if scope.start and scope.end else f"Last {scope.days} days ({span})"


def filename(scope: Scope, kind: str) -> str:
    dates = scope.dates
    return f"dashboard-{dates[0]:%Y-%m-%d}-to-{dates[-1]:%Y-%m-%d}.{kind}"


def _number(value: float | None) -> float | int | None:
    if value is None:
        return None
    return int(value) if float(value).is_integer() else round(value, 1)


def _figure(metric: dict[str, Any]) -> Cell:
    """The value with its unit: 1284 stays a number, 62.5% and 18.5 days read as text."""
    value = _number(metric["value"])
    if value is None:
        return None
    return value if metric["unit"] == "count" else f"{value}{UNITS[metric['unit']]}"


def _change(metric: dict[str, Any]) -> str:
    delta = _number(metric["delta"])
    if delta is None:
        return ""
    sign = "+" if delta > 0 else ""
    return f"{sign}{delta}{DELTA_UNITS[metric['unit']]}"


def _name(user: Any) -> str:
    return user.full_name


def report(scope: Scope, job_description: str | None = None) -> Report:
    """Every widget on the page as a table, in the page's order."""
    window = window_label(scope)
    summary = services.summary(scope)
    attention = services.attention(scope)
    trends = services.trends(scope)
    insights = services.insights(scope)
    pipeline = services.pipeline(scope)
    funnel = services.funnel(scope, job_description)
    interviews = services.interview_insights(scope)
    team = services.team(scope)
    upcoming = services.upcoming_interviews(scope)

    jobs = [job for job in pipeline["jobs"] if job.status in ACTIVE_ROLE_STATUSES]
    jobs.sort(key=lambda job: (-sum(getattr(job, key) for key in ROLE_STAGES), -job.total))
    funnel_job = next((job for job in pipeline["jobs"] if str(job.id) == job_description), None)
    match, offers, outreach, searches = (
        insights["match"],
        insights["offers"],
        insights["outreach"],
        insights["searches"],
    )
    avg_response = _number(offers["avg_response_days"])

    lines = [window]
    if scope.user_ids:
        people = User.objects.filter(pk__in=scope.user_ids).order_by("first_name", "last_name")
        lines.append(f"Job descriptions of {', '.join(_name(user) for user in people)}")
    lines.append(
        f"Exported {scope.now.astimezone(scope.tz):%d %b %Y, %H:%M} by {_name(scope.viewer)}"
    )

    tables = [
        Table(
            "Headline figures",
            ["Figure", "Value", "Change vs previous window", "Covers", "Note"],
            [
                [
                    label,
                    _figure(summary[key]),
                    _change(summary[key]),
                    "Right now" if covers == "now" else window,
                    summary[key]["detail"] or "",
                ]
                for key, label, covers in HEADLINE_FIGURES
            ],
        ),
        Table(
            "Needs attention",
            ["Item", "What it means", "Count"],
            [[label, meaning, attention[key]] for key, label, meaning in ATTENTION_ITEMS],
            note="Right now",
        ),
        Table(
            "Hiring activity",
            ["Day", *(label for _key, label in TREND_SERIES)],
            [
                [point["date"], *(point[key] for key, _label in TREND_SERIES)]
                for point in trends["points"]
            ],
            note=f"{window}, by day",
        ),
        Table(
            "Pipeline health",
            ["Stage", "Candidates", "Avg days in stage", "Over a week"],
            [
                [stage["label"], stage["value"], _number(stage["avg_days"]), stage["stuck"]]
                for stage in insights["stages"]
            ],
            note="Right now",
        ),
        Table(
            "Recruitment funnel",
            ["Stage", "Candidates", "Of previous (%)"],
            [
                [stage["label"], stage["value"], stage["conversion_pct"]]
                for stage in funnel["stages"]
            ],
            note=funnel_job.title if funnel_job else "All roles",
        ),
        Table(
            "Open roles",
            [
                "Role",
                "Department",
                "Status",
                "Openings",
                "Awaiting",
                "Shortlisted",
                "Contacted",
                "Interviewing",
                "Selected",
                "Onboarding",
                "Parked",
                "Total",
            ],
            [
                [
                    job.title,
                    job.department,
                    job.get_status_display(),
                    job.openings,
                    job.awaiting,
                    *(getattr(job, key) for key in ROLE_STAGES),
                    job.parked,
                    job.total,
                ]
                for job in jobs
            ],
            note="Right now, candidates in play per role",
        ),
        Table(
            "Skills in demand",
            ["Skill", "Open roles asking", "Candidates who have it"],
            [[skill["label"], skill["roles"], skill["candidates"]] for skill in insights["skills"]],
        ),
        Table(
            "Match quality",
            ["Match score", "Candidates"],
            [[band["label"], band["value"]] for band in match["bands"]],
            note=(
                f"Right now, {match['scored']} scored candidates"
                + (f", {match['avg_pct']}% on average" if match["avg_pct"] is not None else "")
            ),
        ),
        Table(
            "Experience mix",
            ["Experience", "Candidates"],
            [[band["label"], band["value"]] for band in insights["experience"]],
            note="Right now",
        ),
        Table(
            "AI searches",
            ["Figure", "Value"],
            [
                ["Searches run", searches["runs"]],
                ["Profiles found", searches["found"]],
                ["AI shortlisted", searches["shortlisted"]],
                ["New to the database", searches["new"]],
                [
                    "Average search time (s)",
                    None
                    if searches["avg_duration_ms"] is None
                    else round(searches["avg_duration_ms"] / 1000),
                ],
            ],
            note=window,
        ),
        Table(
            "Candidates by source",
            ["Source", "Candidates"],
            [[source["label"], source["value"]] for source in pipeline["sources"]],
            note="Right now",
        ),
        Table(
            "Departments",
            ["Department", "Open roles", "Openings", "Candidates"],
            [
                [row["label"], row["roles"], row["openings"], row["candidates"]]
                for row in insights["departments"]
            ],
            note="Right now",
        ),
        Table(
            "Offers",
            ["Status", "Offers"],
            [[row["label"], row["value"]] for row in offers["statuses"]],
            note=(
                f"Right now; {offers['responded']} answered in the window"
                + (
                    f", {avg_response} days to answer on average"
                    if avg_response is not None
                    else ""
                )
            ),
        ),
        Table(
            "Outreach",
            ["Channel or outcome", "Logged"],
            [
                *([row["label"], row["value"]] for row in outreach["channels"]),
                *([f"Outcome: {row['label']}", row["value"]] for row in outreach["outcomes"]),
            ],
            note=window,
        ),
        Table(
            "Interviews",
            ["Figure", "Value"],
            [
                ["Held", interviews["total"]],
                ["Completed", interviews["completed"]],
                ["Cancelled", interviews["cancelled"]],
                ["No-show", interviews["no_show"]],
                ["Average score", interviews["avg_score"]],
                ["Upcoming (right now)", interviews["upcoming"]],
                ["Feedback owed (right now)", interviews["feedback_pending"]],
                *(
                    [f"Recommendation: {row['label']}", row["value"]]
                    for row in interviews["recommendations"]
                ),
            ],
            note=window,
        ),
        Table(
            "Interviewer load",
            ["Interviewer", "Held", "Completed", "Avg score"],
            [
                [_name(row["user"]), row["total"], row["completed"], row["avg_score"]]
                for row in interviews["interviewers"]
            ],
            note=window,
        ),
        Table(
            "When the team works",
            ["Hour", *WEEKDAYS],
            [[f"{hour:02d}:00", *(row[hour] for row in insights["heatmap"])] for hour in range(24)],
            note=f"{window}, actions by hour of the day in your time zone",
        ),
        Table(
            "Team activity",
            ["Person", "Roles", "Sourcing", "Outreach", "Interviews", "Closing", "Total"],
            [
                [
                    _name(member["user"]),
                    member["roles"],
                    member["sourcing"],
                    member["outreach"],
                    member["interviews"],
                    member["closing"],
                    member["total"],
                ]
                for member in team
            ],
            note=f"{window}, across everything you can see",
        ),
        Table(
            "Upcoming interviews",
            ["When", "Round", "Candidate", "Role", "Interviewer"],
            [
                [
                    interview.scheduled_at.astimezone(scope.tz).replace(tzinfo=None),
                    interview.get_round_display(),
                    interview.application.candidate.full_name,
                    interview.application.job_description.title,
                    _name(interview.interviewer),
                ]
                for interview in upcoming
            ],
            note="The next five, in your time zone",
        ),
    ]
    return Report("Dashboard", lines, tables)


# ------------------------------------------------------------------ writers


def _text(cell: Cell) -> str:
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return f"{cell:%a, %d %b %Y %H:%M}"
    if isinstance(cell, date):
        return f"{cell:%a, %d %b %Y}"
    return str(cell)


def to_csv(report: Report) -> bytes:
    """One file, each table under its title, with a byte-order mark so Excel reads UTF-8."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([report.title])
    writer.writerows([line] for line in report.lines)
    for table in report.tables:
        writer.writerow([])
        writer.writerow([table.title, table.note] if table.note else [table.title])
        writer.writerow(table.columns)
        writer.writerows([_text(cell) for cell in row] for row in table.rows)
    return ("﻿" + buffer.getvalue()).encode("utf-8")


HEADER_FILL = PatternFill("solid", fgColor="E8EDF5")


def to_xlsx(report: Report) -> bytes:
    """One sheet per table, a bold frozen header row and dates as real dates."""
    book = Workbook()
    about = book.active
    about.title = "About"
    about.append([report.title])
    about["A1"].font = Font(bold=True, size=14)
    for line in report.lines:
        about.append([line])
    about.column_dimensions["A"].width = 80

    for table in report.tables:
        sheet = book.create_sheet(table.title)
        if table.note:
            sheet.append([table.note])
            sheet["A1"].font = Font(italic=True, color="666666")
        sheet.append(table.columns)
        header = sheet.max_row
        for cell in sheet[header]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        for row in table.rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, datetime):
                    cell.number_format = "ddd, d mmm yyyy hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "ddd, d mmm yyyy"
        sheet.freeze_panes = sheet.cell(row=header + 1, column=1)
        for index, column in enumerate(table.columns, start=1):
            longest = max([len(column), *(len(_text(row[index - 1])) for row in table.rows)])
            sheet.column_dimensions[get_column_letter(index)].width = min(longest, 60) + 3

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


PDF_CSS = """
body { font-family: sans-serif; font-size: 8.5pt; color: #1a1a1a; }
h1 { font-size: 18pt; margin: 0; }
h2 { font-size: 11pt; margin: 16pt 0 0 0; }
p { margin: 2pt 0 0 0; color: #555555; }
table { border-collapse: collapse; margin-top: 5pt; }
th, td { border: 1px solid #cfd4dc; padding: 2pt 6pt; text-align: left; }
th { font-weight: bold; border-bottom: 2px solid #9aa3b2; }
"""
PAGE_MARGIN = 36
# Story cannot paginate a table taller than a page once it has started near the foot
# of one (it lays the same rows out again for ever), so long tables go in as a run of
# short tables, each with the header row, that always fit on a page.
PDF_ROWS_PER_TABLE = 15


def to_pdf(report: Report) -> bytes:
    """Landscape A4 pages (the open roles table is twelve columns wide) laid out from
    HTML by PyMuPDF's Story engine. Header cells carry no background fill: Story
    repaints fills at the same spot on every later page."""
    parts = [
        f"<h1>{escape(report.title)}</h1>",
        *(f"<p>{escape(line)}</p>" for line in report.lines),
    ]
    for table in report.tables:
        parts.append(f"<h2>{escape(table.title)}</h2>")
        if table.note:
            parts.append(f"<p>{escape(table.note)}</p>")
        head = "".join(f"<th>{escape(column)}</th>" for column in table.columns)
        for first in range(0, max(len(table.rows), 1), PDF_ROWS_PER_TABLE):
            body = "".join(
                "<tr>" + "".join(f"<td>{escape(_text(cell))}</td>" for cell in row) + "</tr>"
                for row in table.rows[first : first + PDF_ROWS_PER_TABLE]
            )
            parts.append(f"<table><tr>{head}</tr>{body}</table>")

    story = pymupdf.Story(html="".join(parts), user_css=PDF_CSS)
    page = pymupdf.paper_rect("a4-l")
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(page + (PAGE_MARGIN, PAGE_MARGIN, -PAGE_MARGIN, -PAGE_MARGIN))
        story.draw(device)
        writer.end_page()
    writer.close()
    return buffer.getvalue()


WRITERS = {
    "csv": ("text/csv; charset=utf-8", to_csv),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx),
    "pdf": ("application/pdf", to_pdf),
}
